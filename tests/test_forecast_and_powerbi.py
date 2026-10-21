"""Forecast, model reliability check, Power BI files and automatic refresh."""
import datetime
import io
import os
from pathlib import Path
import re
import time
import zipfile

import numpy as np
import pandas as pd
import pytest

from src.agent import SalesAgent
from src.data import clean_data
from src.export import quality_report_frame
from src.forecast import forecast_monthly, forecast_table, monthly_summary
from src.model import cross_validate_time, reliability_verdict, score_dataframe, train_and_evaluate
from src.pipeline import list_input_files, run_pipeline, watch
import src.powerbi as powerbi_module
from src.powerbi import POWERBI_DIR, build_powerbi_files, zip_bytes
from src.rag import SalesRAG
from src.security import prepare_table

ROOT = Path(__file__).parent.parent


@pytest.fixture(scope="module")
def demo():
    df, report = clean_data(prepare_table(pd.read_csv(ROOT / "data" / "sample_superstore.csv")))
    return df, report


def monthly_sales_table(months=48, noise=0.03, seed=0, last_day=28):
    """Orders every 3 days; monthly totals follow trend + seasonality, so the true future is known."""
    rng = np.random.default_rng(seed)
    t = np.arange(months + 3)
    truth = 10000 + 150 * t + 3000 * np.sin(2 * np.pi * t / 12)
    rows = []
    for i, period in enumerate(pd.period_range("2021-01", periods=months, freq="M")):
        for day in range(0, last_day, 3):
            rows.append({"Order Date": period.to_timestamp() + pd.Timedelta(days=day),
                         "Sales": truth[i] * rng.normal(1, noise) / 10, "Profit": truth[i] / 100})
    return pd.DataFrame(rows), truth[months:]


# ----------------------------------------------------------------------------- forecast
def test_forecast_is_close_to_a_known_future():
    table, truth = monthly_sales_table()
    result = forecast_monthly(table, "Sales", 3)
    assert result["status"] == "ok" and result["method"] in ("trend_seasonal", "seasonal_level")
    assert np.allclose(result["forecast"]["forecast"], truth, rtol=0.08)
    assert (result["forecast"]["lower"] <= result["forecast"]["forecast"]).all()
    assert (result["forecast"]["forecast"] <= result["forecast"]["upper"]).all()
    assert result["wape"] <= result["baseline_wape"] and result["verdict"] == "good"


def test_forecast_says_so_when_history_is_too_short():
    table, _ = monthly_sales_table(months=7)
    result = forecast_monthly(table, "Sales", 1)
    assert result["status"] == "not_enough" and result["n_months"] == 7


def test_incomplete_last_month_is_not_used_for_learning():
    table, _ = monthly_sales_table(months=30)
    last = table["Order Date"].dt.to_period("M").max()
    table = table[~((table["Order Date"].dt.to_period("M") == last) & (table["Order Date"].dt.day > 10))]
    result = forecast_monthly(table, "Sales", 1)
    assert result["notes"]["dropped_last"] == last
    assert result["forecast"].index[0] == last              # the forecast starts at the unfinished month
    assert result["history"].index[-1] == last - 1


def test_forecast_never_claims_reliability_for_wild_data():
    rng = np.random.default_rng(5)
    table, _ = monthly_sales_table(months=36)
    table["Sales"] = rng.lognormal(4, 1.6, len(table))         # huge random swings
    assert forecast_monthly(table, "Sales", 1)["verdict"] != "good"


def test_profit_forecast_can_be_negative_but_sales_cannot():
    table, _ = monthly_sales_table(months=24)
    table["Profit"] = -abs(table["Profit"])
    assert (forecast_monthly(table, "Profit", 2)["forecast"]["forecast"] < 0).all()
    table["Sales"] = 0.0
    assert (forecast_monthly(table, "Sales", 2)["forecast"]["forecast"] >= 0).all()


def test_forecast_tables_for_power_bi(demo):
    df, _ = demo
    flat = forecast_table(forecast_monthly(df, "Sales", 3))
    assert list(flat.columns) == ["Metric", "Month", "Forecast", "Lower", "Upper", "Method", "Backtest WAPE", "Baseline WAPE", "Reliability"]
    assert len(flat) == 3 and flat["Month"].is_monotonic_increasing
    summary = monthly_summary(df)
    assert summary["Month Complete"].isin([0, 1]).all() and len(summary) == 48


def test_assistant_answers_forecast_questions_in_both_languages(demo):
    df, _ = demo
    agent = SalesAgent(df, SalesRAG(df))
    english = agent.run("What is the sales forecast for next month?")
    arabic = agent.run("توقع مبيعات الشهر الجاي")
    assert english["tool"] == arabic["tool"] == "forecast"
    assert "Forecast for" in english["answer"] and "توقع" in arabic["answer"]


# ----------------------------------------------------------------------------- reliability check
def test_reliability_is_strong_on_learnable_data_and_weak_on_noise(demo):
    df, _ = demo
    assert reliability_verdict(cross_validate_time(df))["verdict"] == "strong"

    rng = np.random.default_rng(1)
    noise = df.sample(3000, random_state=1).sort_values("Order Date").reset_index(drop=True)
    noise["Profitable"] = rng.integers(0, 2, len(noise))      # the outcome has nothing to do with the features
    verdict = reliability_verdict(cross_validate_time(noise))
    assert verdict["verdict"] in ("weak", "not_enough")


def test_reliability_needs_at_least_two_usable_periods():
    assert reliability_verdict(pd.DataFrame())["verdict"] == "not_enough"


def test_scored_data_marks_rows_used_for_training(demo):
    df, _ = demo
    model, features, metrics = train_and_evaluate(df)
    scored = score_dataframe(model, df, metrics["train_index"])
    assert scored["Profit Probability"].between(0, 1).all()
    assert set(scored["Risk Level"]) <= {"Low", "Medium", "High"}
    assert scored["In Training Set"].sum() == metrics["train_size"]
    unseen = scored[scored["In Training Set"] == 0]
    assert (unseen["Predicted Profitable"] == unseen["Profitable"]).mean() == pytest.approx(metrics["accuracy"], abs=0.002)
    assert (scored.loc[scored["Risk Level"] == "High", "Profit Probability"] <= 0.4).all()


# ----------------------------------------------------------------------------- power bi pack
def test_power_bi_pack_contents(demo):
    df, report = demo
    bundle = train_and_evaluate(df)
    files = build_powerbi_files(df, report, bundle)
    assert {"cleaned_sales.csv", "sales_scored.csv", "forecast.csv", "monthly_summary.csv", "quality_report.csv",
            "model_metrics.csv", "refresh_status.csv", "measures.dax", "load_sales.pq", "README_POWERBI.md",
            "power_bi_data.xlsx", "STEP_BY_STEP.md", "power_bi_dashboard.zip"} <= set(files)
    assert "sales_scored.csv" not in build_powerbi_files(df, report, None)     # no model, no predictions
    assert zipfile.ZipFile(io.BytesIO(zip_bytes(files))).testzip() is None


def test_power_bi_columns_stay_english_for_an_arabic_file():
    arabic = pd.read_csv(ROOT / "data" / "sample_arabic_sales.csv", encoding="utf-8-sig").head(600)
    df, report = clean_data(prepare_table(arabic))
    scored = pd.read_csv(io.BytesIO(build_powerbi_files(df, report, train_and_evaluate(df))["sales_scored.csv"]), encoding="utf-8-sig")
    assert {"Order Date", "Sales", "Profit", "Category", "Profit Probability", "Risk Level"} <= set(scored.columns)


def test_dax_measures_only_use_columns_that_the_files_contain(demo):
    """The DAX file must match the exported files, otherwise Power BI shows errors."""
    df, report = demo
    files = build_powerbi_files(df, report, train_and_evaluate(df))
    sales_columns = set(pd.read_csv(io.BytesIO(files["sales_scored.csv"]), encoding="utf-8-sig", nrows=1).columns)
    forecast_columns = set(pd.read_csv(io.BytesIO(files["forecast.csv"]), encoding="utf-8-sig").columns)
    dax = (POWERBI_DIR / "measures.dax").read_text(encoding="utf-8")
    dax = re.sub(r"--.*", "", dax)                                          # ignore comments
    assert set(re.findall(r"Sales\[([^\]]+)\]", dax)) <= sales_columns | {"Discount Band"}
    assert set(re.findall(r"Forecast\[([^\]]+)\]", dax)) <= forecast_columns
    for measure in ["Total Sales", "Total Profit", "Profit Margin %", "Loss Rate %", "Avg Discount", "Sales LY"]:
        assert f"{measure} =" in dax                                        # the measures the user asked for
    assert dax.count("(") == dax.count(")")                                  # balanced brackets


# ----------------------------------------------------------------------------- one Excel file for Power BI
def test_single_excel_workbook_has_every_table_with_the_right_types(demo):
    from openpyxl import load_workbook
    df, report = demo
    tricky = df.head(2000).copy()
    tricky.loc[tricky.index[0], "Category"] = "=1+1"                      # must never become a formula
    tricky.loc[tricky.index[1], "Customer Name"] = "bad\x07name"           # a control character Excel rejects
    files = build_powerbi_files(tricky, report, train_and_evaluate(tricky))
    workbook = load_workbook(io.BytesIO(files["power_bi_data.xlsx"]))
    assert {"Sales", "Forecast", "Monthly", "Quality", "ModelMetrics", "Status", "README"} <= set(workbook.sheetnames)
    for name in ["Sales", "Forecast", "Monthly", "Quality", "ModelMetrics", "Status"]:
        assert list(workbook[name].tables) == [name]                      # an Excel Table named like the sheet
    sales = workbook["Sales"]
    csv = pd.read_csv(io.BytesIO(files["sales_scored.csv"]), encoding="utf-8-sig")
    assert [c.value for c in sales[1]] == list(csv.columns)               # same columns as the CSV
    assert sales.max_row - 1 == len(csv)
    headers = [c.value for c in sales[1]]
    assert isinstance(sales.cell(row=2, column=headers.index("Order Date") + 1).value, datetime.datetime)  # a real date
    assert not any(cell.data_type == "f" for row in sales.iter_rows() for cell in row)   # no formulas anywhere
    assert "bad\x07name" not in {c.value for row in sales.iter_rows() for c in row}


def test_workbook_is_skipped_for_very_large_data_but_csv_files_stay(demo, monkeypatch):
    df, report = demo
    monkeypatch.setattr(powerbi_module, "EXCEL_ROW_LIMIT", 100)
    files = powerbi_module.build_powerbi_files(df.head(500), report)
    assert "power_bi_data.xlsx" not in files and "cleaned_sales.csv" in files


def test_step_by_step_guide_matches_the_dax_file_and_the_workbook(demo):
    guide = (POWERBI_DIR / "STEP_BY_STEP.md").read_text(encoding="utf-8")
    dax = re.sub(r"--.*", "", (POWERBI_DIR / "measures.dax").read_text(encoding="utf-8"))
    section = guide.split("## 3)")[1].split("لتنسيق النسب")[0]            # only the ordered list of measures
    listed = [token for line in section.splitlines() if re.match(r"^\d\. ", line) for token in re.findall(r"`([^`]+)`", line)]
    assert len(listed) >= 25
    for name in listed:
        assert re.search(rf"(^|\n)\s*{re.escape(name)}\s*=", dax), f"{name} is in the guide but not in measures.dax"
    for table in ["Sales", "Forecast", "Monthly", "Quality", "ModelMetrics", "Status"]:
        assert table in guide
    df, report = demo
    quality = quality_report_frame(report, "en")["Metric"].tolist()
    assert "Total sales after cleaning" in quality and "Total sales after cleaning" in guide   # the guide's check step


def test_app_offers_the_single_excel_file(monkeypatch):
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    [b.click() for b in at.button if "Power BI" in b.label or "حزمة" in b.label]
    at.run()
    pack, count, workbook, dashboard = at.session_state["pbi_pack"]
    assert not at.exception and workbook is not None and dashboard is not None and count >= 9
    assert zipfile.ZipFile(io.BytesIO(pack)).namelist().count("power_bi_data.xlsx") == 1


# ----------------------------------------------------------------------------- automatic refresh
def sales_file(path, rows=700):
    pd.read_csv(ROOT / "data" / "sample_arabic_sales.csv", encoding="utf-8-sig").head(rows).to_csv(path, index=False, encoding="utf-8-sig")


def test_refresh_writes_all_files_safely(tmp_path):
    inbox, exports = tmp_path / "inbox", tmp_path / "exports"
    inbox.mkdir()
    assert run_pipeline(inbox, exports)["status"] == "NO_FILES"
    sales_file(inbox / "sales.csv")
    result = run_pipeline(inbox, exports, reliability=False)
    assert result["status"] == "OK" and (exports / "sales_scored.csv").exists()
    assert (exports / "power_bi_data.xlsx").exists() and (exports / "power_bi_dashboard.zip").exists()
    assert (exports / "profit_model.joblib").exists() and (exports / "use_model.py").exists()    # the model for their own system
    assert list(exports.glob("*.tmp")) == []
    status = pd.read_csv(exports / "refresh_status.csv")
    assert status.loc[0, "Status"] == "OK" and status.loc[0, "Source"] == "sales.csv"


def test_failed_refresh_keeps_the_previous_exports(tmp_path):
    inbox, exports = tmp_path / "inbox", tmp_path / "exports"
    inbox.mkdir()
    sales_file(inbox / "sales.csv")
    run_pipeline(inbox, exports, reliability=False)
    before = (exports / "cleaned_sales.csv").read_bytes()
    (inbox / "broken.csv").write_text("a,b\n1,2\n")
    os.utime(inbox / "broken.csv", (time.time() + 10, time.time() + 10))     # the newest file
    assert run_pipeline(inbox, exports)["status"] == "FAILED"
    assert (exports / "cleaned_sales.csv").read_bytes() == before
    assert pd.read_csv(exports / "refresh_status.csv").loc[0, "Status"] == "FAILED"


def test_small_file_still_exports_everything_except_predictions(tmp_path):
    inbox, exports = tmp_path / "inbox", tmp_path / "exports"
    inbox.mkdir()
    sales_file(inbox / "tiny.csv", rows=20)
    result = run_pipeline(inbox, exports)
    assert result["status"] == "OK" and "At least 30" in result["message"]
    assert (exports / "cleaned_sales.csv").exists() and not (exports / "sales_scored.csv").exists()


def test_combine_merges_files_and_removes_overlap(tmp_path):
    inbox, exports = tmp_path / "inbox", tmp_path / "exports"
    inbox.mkdir()
    full = pd.read_csv(ROOT / "data" / "sample_arabic_sales.csv", encoding="utf-8-sig")
    full.iloc[:800].to_csv(inbox / "a.csv", index=False, encoding="utf-8-sig")
    full.iloc[500:1500].to_csv(inbox / "b.csv", index=False, encoding="utf-8-sig")
    assert run_pipeline(inbox, exports, combine=True, reliability=False)["rows"] == 1500 - full.iloc[:1500].duplicated().sum()
    assert len(list_input_files(inbox)) == 2


def test_watch_refreshes_only_when_the_folder_changes(tmp_path):
    inbox, exports = tmp_path / "inbox", tmp_path / "exports"
    inbox.mkdir()
    sales_file(inbox / "sales.csv", rows=300)
    messages = []
    watch(inbox, exports, interval=0.01, settle=0, max_cycles=3, report=messages.append, reliability=False)
    assert len(messages) == 1 and messages[0].split()[1] == "OK"


def test_excel_lock_files_in_the_inbox_are_ignored(tmp_path):
    (tmp_path / "~$sales.xlsx").write_bytes(b"lock")
    sales_file(tmp_path / "sales.csv", rows=50)
    assert [p.name for p in list_input_files(tmp_path)] == ["sales.csv"]


def test_forecast_says_when_it_is_only_last_months_value(demo):
    df, _ = demo
    result = forecast_monthly(df, "Profit", 1)
    assert result["method"] == "naive"                                      # nothing beat "repeat last month"
    assert "equals last month" in SalesAgent(df, SalesRAG(df)).run("profit forecast next month")["answer"]


def test_forecast_accepts_the_names_shown_on_screen_in_both_languages(demo):
    """A KeyError 'Column not found: المبيعات' happened when the screen label reached the forecast code."""
    from src.forecast import normalize_metric
    df, _ = demo
    assert [normalize_metric(v) for v in ("Sales", "Profit", "المبيعات", "الربح")] == ["Sales", "Profit", "Sales", "Profit"]
    assert forecast_monthly(df, "المبيعات", 1)["metric"] == "Sales"
    assert forecast_monthly(df, "الربح", 1)["metric"] == "Profit"
    with pytest.raises(ValueError, match="no 'Profit' column"):
        forecast_monthly(df.drop(columns=["Profit"]), "Profit", 1)


def test_app_survives_when_the_widgets_return_the_text_shown_on_screen(monkeypatch):
    """Even if Streamlit returned the Arabic label instead of the option, the page must not crash."""
    from streamlit.delta_generator import DeltaGenerator
    from streamlit.testing.v1 import AppTest
    original = DeltaGenerator.radio

    def radio_returning_labels(self, label, options=(), *args, **kwargs):
        value = original(self, label, options, *args, **kwargs)
        if kwargs.get("key") == "fc_metric":
            return "المبيعات"
        if kwargs.get("key") == "demo_choice":
            return "بيانات إنجليزية (Superstore)"
        return value

    monkeypatch.setattr(DeltaGenerator, "radio", radio_returning_labels)
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    assert not at.exception
    assert any("2027" in m.label for m in at.metric)                  # the forecast was calculated


def test_a_failing_forecast_shows_a_message_and_the_rest_of_the_page_still_works(monkeypatch):
    import src.forecast as forecast_module
    from streamlit.testing.v1 import AppTest

    def broken(*args, **kwargs):
        raise RuntimeError("something went wrong inside the forecast")

    monkeypatch.setattr(forecast_module, "forecast_monthly", broken)
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300).run()
    assert not at.exception                                              # the page did not crash
    assert any("something went wrong" in e.value for e in at.error)      # the user is told
    assert len(at.tabs) == 6 and at.metric                               # the other tabs are still there
