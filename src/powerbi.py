"""Builds the files for Power BI (used by the app's download button and by refresh_exports.py)."""
import io
from pathlib import Path
import re
import zipfile

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .export import clean_csv_bytes, protect_text, quality_report_frame
from .forecast import forecast_monthly, forecast_table, monthly_summary
from .model import score_dataframe
from .pbip import build_project

POWERBI_DIR = Path(__file__).resolve().parent.parent / "powerbi"
DOCUMENT_FILES = ["measures.dax", "load_sales.pq", "README_POWERBI.md", "STEP_BY_STEP.md"]
WORKBOOK_NAME = "power_bi_data.xlsx"
DASHBOARD_ZIP = "power_bi_dashboard.zip"
EXCEL_ROW_LIMIT = 100_000   # writing Excel is slow (about 4 s per 10,000 rows); above this only the CSV files are made
ILLEGAL_EXCEL_CHARACTERS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def metrics_frame(metrics, reliability=None):
    """Model quality as a small table: Metric, Value, Note."""
    rows = [
        ("Accuracy", metrics["accuracy"], "tested on the most recent rows"),
        ("Baseline Accuracy", metrics["baseline_accuracy"], "best constant guess; the model must beat this"),
        ("Precision", metrics["precision"], ""),
        ("Recall", metrics["recall"], ""),
        ("F1", metrics["f1"], ""),
        ("ROC-AUC", metrics["roc_auc"], ""),
        ("Train Rows", metrics["train_size"], ""),
        ("Test Rows", metrics["test_size"], metrics["split_strategy"]),
    ]
    if reliability and reliability.get("verdict") != "not_enough":
        rows += [
            ("CV Periods", reliability["folds"], "time-based cross-validation"),
            ("CV Accuracy", reliability["accuracy"], ""),
            ("CV Baseline", reliability["baseline"], ""),
            ("CV Gain Over Baseline", reliability["gain"], ""),
            ("CV ROC-AUC", reliability["roc_auc"], ""),
            ("CV Spread", reliability["spread"], "std of accuracy between periods"),
            ("Reliability Verdict", float("nan"), reliability["verdict"]),
        ]
    table = pd.DataFrame(rows, columns=["Metric", "Value", "Note"])
    table["Value"] = table["Value"].astype(float).round(4)
    return table


def status_frame(source, rows_loaded, rows_clean, status, message=""):
    return pd.DataFrame([{
        "Last Refresh": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Status": status,
        "Source": source,
        "Rows Loaded": rows_loaded,
        "Rows After Cleaning": rows_clean,
        "Message": message,
    }])


def excel_safe(frame):
    """Text that Excel could run as a formula gets a leading apostrophe; control characters Excel rejects are removed."""
    safe = frame.copy()
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        safe[column] = safe[column].map(
            lambda value: protect_text(ILLEGAL_EXCEL_CHARACTERS.sub("", value)) if isinstance(value, str) else value
        )
    return safe


def workbook_bytes(tables, readme_lines):
    """ONE Excel file for Power BI. Every table is an Excel Table with a fixed name (Sales, Forecast, ...),
    so Power BI lists them by name, and the names match the DAX measures."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as writer:
        for name, frame in tables.items():
            excel_safe(frame).to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            last_column = get_column_letter(max(len(frame.columns), 1))
            table = Table(displayName=name, ref=f"A1:{last_column}{max(len(frame) + 1, 2)}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
            for index, column in enumerate(frame.columns, start=1):
                if pd.api.types.is_datetime64_any_dtype(frame[column]):   # show dates without 00:00:00
                    for (cell,) in sheet.iter_rows(min_row=2, max_row=len(frame) + 1, min_col=index, max_col=index):
                        cell.number_format = "yyyy-mm-dd"
                longest = max([len(str(column))] + [len(str(v)) for v in frame[column].head(200)])
                sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 40)
        guide = writer.book.create_sheet("README")
        for row, line in enumerate(readme_lines, start=1):
            guide.cell(row=row, column=1, value=line)
        guide["A1"].font = Font(bold=True, size=14)
        guide.column_dimensions["A"].width = 110
    return buffer.getvalue()


WORKBOOK_README = [
    "Sales Intelligence - data for Power BI | بيانات Power BI",
    "",
    "Power BI Desktop: Home > Get data > Excel workbook > choose this file > tick ALL the tables > Load.",
    "Tables: Sales (cleaned data, with profit probability when the model was trained), Forecast, Monthly, Quality, ModelMetrics, Status.",
    "The table names match the DAX measures in measures.dax. Follow STEP_BY_STEP.md for the Date table, the measures and the visuals.",
    "",
    "Power BI Desktop: الصفحة الرئيسية > Get data > Excel workbook > اختر هذا الملف > علّم كل الجداول > Load.",
    "أسماء الجداول ثابتة: Sales و Forecast و Monthly و Quality و ModelMetrics و Status، وتطابق معادلات DAX.",
    "Rows with 'In Training Set' = 1 were used to train the model, so their probabilities are optimistic.",
]


def dashboard_zip(project_files, workbook):
    """The Power BI project (.pbip) with the data file next to it. Unzip it into C:\\SalesIntelligence\\ ."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path, text in project_files.items():
            archive.writestr(path, text)
        archive.writestr(WORKBOOK_NAME, workbook)
    return buffer.getvalue()


def build_powerbi_files(df, report, model_bundle=None, reliability=None, horizon=3, source="app", lang="en"):
    """Returns {file name: bytes}. Predictions are included only when a trained model is given."""
    sales = df
    tables_csv = {"cleaned_sales.csv": df}
    metrics_table = None
    if model_bundle is not None:
        model, _features, metrics = model_bundle
        sales = score_dataframe(model, df, metrics["train_index"])
        tables_csv["sales_scored.csv"] = sales
        metrics_table = metrics_frame(metrics, reliability)
        tables_csv["model_metrics.csv"] = metrics_table
    forecast = pd.concat([forecast_table(forecast_monthly(df, metric, horizon)) for metric in ("Sales", "Profit")], ignore_index=True)
    monthly = monthly_summary(df)
    quality = quality_report_frame(report, "en")
    status = status_frame(source, report["rows_before"], len(df), "OK")
    tables_csv.update({"forecast.csv": forecast, "monthly_summary.csv": monthly,
                       "quality_report.csv": quality, "refresh_status.csv": status})
    files = {name: clean_csv_bytes(table) for name, table in tables_csv.items()}

    if len(sales) <= EXCEL_ROW_LIMIT:   # the single Excel file with every table
        tables = {"Sales": sales, "Forecast": forecast, "Monthly": monthly, "Quality": quality}
        if metrics_table is not None:
            tables["ModelMetrics"] = metrics_table
        tables["Status"] = status
        files[WORKBOOK_NAME] = workbook_bytes(tables, WORKBOOK_README)
        # The ready dashboard (a Power BI project). It reads the tables of the workbook above.
        project = build_project(sales, model_bundle is not None, len(forecast) > 0, lang)
        files[DASHBOARD_ZIP] = dashboard_zip(project, files[WORKBOOK_NAME])

    for name in DOCUMENT_FILES:
        path = POWERBI_DIR / name
        if path.exists():
            files[name] = path.read_bytes()
    return files


def zip_bytes(files):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()
